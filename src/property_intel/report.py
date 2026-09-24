from collections import Counter
from datetime import date
from pathlib import Path
from statistics import median

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from sqlalchemy import select
from sqlalchemy.orm import Session

from .analytics import comparables, overview
from .db import Property, RawObservation, Transaction

NAVY = "18252C"
TEAL = "196C68"
PALE = "EAF2F0"


def export_report(
    session: Session,
    path: Path = Path("data/exports/NZ_Commercial_Property_Market_Report.xlsx"),
    subject: Property | None = None,
) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    summary = book.active
    assert summary is not None
    summary.title = "01 Executive Summary"
    transactions = book.create_sheet("02 Transactions")
    comps = book.create_sheet("03 Comparable Sales")
    sectors = book.create_sheet("04 Sector Analysis")
    quality = book.create_sheet("05 Data Quality")
    stats = overview(session)
    summary.append(["PROPERTY INTEL  /  MARKET REPORT"])
    summary.append(
        [
            "DEMO DATA — SYNTHETIC"
            if session.scalar(select(RawObservation.id).where(RawObservation.import_job_id == 1))
            else "LOCAL DATA"
        ]
    )
    summary.append(["Reporting period", f"Through {date.today():%d %B %Y}"])
    for label, value in (
        ("Properties tracked", stats["properties"]),
        ("Transactions", stats["transactions"]),
        ("Tracked transaction value", stats["value"]),
        ("Median sale price", stats["median_price"]),
        ("Median yield", stats["median_yield"]),
        ("Yield sample size", stats["yield_count"]),
        ("Records requiring review", stats["review"]),
    ):
        summary.append([label, value])
    summary["B6"].number_format = summary["B7"].number_format = '"$"#,##0'
    summary["B8"].number_format = "0.00%"
    summary.append([])
    summary.append(["Year", "Transactions"])
    for year, count in sorted(stats["year_counts"].items()):
        summary.append([year, count])
    chart = BarChart()
    chart.title = "Transactions by year"
    chart.add_data(
        Reference(summary, min_col=2, min_row=12, max_row=11 + len(stats["year_counts"])),
        titles_from_data=True,
    )
    chart.set_categories(
        Reference(summary, min_col=1, min_row=12, max_row=11 + len(stats["year_counts"]))
    )
    summary.add_chart(chart, "D4")
    transactions.append(
        [
            "Date",
            "Address",
            "City",
            "Region",
            "Sector",
            "Sale Price",
            "Land Area m²",
            "Building Area m²",
            "$/Land m²",
            "$/Building m²",
            "Yield",
            "Buyer",
            "Seller",
            "Source",
        ]
    )
    sales = session.scalars(select(Transaction).order_by(Transaction.transaction_date.desc())).all()
    for sale in sales:
        p = sale.property
        transactions.append(
            [
                sale.transaction_date,
                p.canonical_address,
                p.city,
                p.region,
                p.sector,
                sale.sale_price,
                sale.land_area_m2,
                sale.building_area_m2,
                sale.price_per_land_m2,
                sale.price_per_building_m2,
                sale.reported_yield,
                sale.buyer_name,
                sale.seller_name,
                sale.source.name if sale.source else None,
            ]
        )
    transactions.freeze_panes = "C2"
    transactions.auto_filter.ref = transactions.dimensions
    if sales:
        excel_table = Table(displayName="TransactionSales", ref=transactions.dimensions)
        excel_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        transactions.add_table(excel_table)
    for row in transactions.iter_rows(min_row=2):
        row[0].number_format = "dd mmm yyyy"
        for i in (5, 8, 9):
            row[i].number_format = '"$"#,##0'
        row[10].number_format = "0.00%"
    comps.append(
        [
            "SUBJECT PROPERTY",
            subject.canonical_address
            if subject
            else "Select a property in the TUI to populate comparables",
        ]
    )
    if subject:
        comps.append(["Sector", subject.sector])
        comps.append(["Land area m²", subject.land_area_m2])
        comps.append(["Building area m²", subject.building_area_m2])
    comps.append([])
    comps.append(
        [
            "Rank",
            "Address",
            "Distance km",
            "Sale Date",
            "Sale Price",
            "Land Area m²",
            "Building Area m²",
            "Yield",
            "Similarity",
            "Why matched",
        ]
    )
    if subject:
        for rank, (sale, score, evidence) in enumerate(comparables(session, subject), 1):
            why = f"{'Same' if evidence['same_sector'] else 'Different'} sector; land {evidence['land_difference']:.0%} difference; {evidence['age_years']:.1f} years old"
            comps.append(
                [
                    rank,
                    sale.property.canonical_address,
                    evidence["distance_km"],
                    sale.transaction_date,
                    sale.sale_price,
                    sale.land_area_m2,
                    sale.building_area_m2,
                    sale.reported_yield,
                    score,
                    why,
                ]
            )
    comps.freeze_panes = "C6"
    comps.auto_filter.ref = f"A5:J{comps.max_row}"
    for cell in comps[5]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color="FFFFFF", bold=True)
    for row in comps.iter_rows(min_row=6):
        row[3].number_format = "dd mmm yyyy"
        row[4].number_format = '"$"#,##0'
        row[7].number_format = "0.00%"
        row[8].number_format = "0%"
    sector_properties = Counter(p.sector for p in session.scalars(select(Property)))
    grouped: dict[str, list[Transaction]] = {}
    for sale in sales:
        grouped.setdefault(sale.property.sector, []).append(sale)
    sectors.append(
        [
            "Sector",
            "Properties",
            "Transactions",
            "Transaction Value",
            "Median Sale Price",
            "Median Yield",
            "Median $/Land m²",
        ]
    )
    for sector, group in sorted(grouped.items()):
        yields = [t.reported_yield for t in group if t.reported_yield is not None]
        units = [t.price_per_land_m2 for t in group if t.price_per_land_m2 is not None]
        sectors.append(
            [
                sector,
                sector_properties[sector],
                len(group),
                sum(t.sale_price for t in group),
                median(t.sale_price for t in group),
                median(yields) if yields else None,
                median(units) if units else None,
            ]
        )
    sector_chart = BarChart()
    sector_chart.title = "Transactions by sector"
    sector_chart.add_data(
        Reference(sectors, min_col=3, min_row=1, max_row=sectors.max_row), titles_from_data=True
    )
    sector_chart.set_categories(Reference(sectors, min_col=1, min_row=2, max_row=sectors.max_row))
    sectors.add_chart(sector_chart, "I2")
    sectors.freeze_panes = "B2"
    sectors.auto_filter.ref = sectors.dimensions
    for row in sectors.iter_rows(min_row=2):
        for cell in (row[3], row[4], row[6]):
            cell.number_format = '"$"#,##0'
        row[5].number_format = "0.00%"
    props = session.scalars(select(Property)).all()
    observations = session.scalars(select(RawObservation)).all()
    quality.append(["Data quality indicator", "Count"])
    for label, count in (
        ("Potential duplicates", sum(o.status == "REVIEW" for o in observations)),
        (
            "Validation warnings",
            sum(bool(o.issues_json and "warning" in o.issues_json) for o in observations),
        ),
        ("Missing owner", sum(not p.owner_name for p in props)),
        ("Missing tenant", sum(not p.tenant_name for p in props)),
        (
            "Suspicious transactions",
            sum(t.reported_yield is not None and t.reported_yield > 0.2 for t in sales),
        ),
        ("Import errors", sum(o.status == "REJECTED" for o in observations)),
    ):
        quality.append([label, count])
    quality.conditional_formatting.add(
        "B2:B7",
        CellIsRule(
            operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor="FFF0D1")
        ),
    )
    for sheet in book:
        sheet.sheet_view.showGridLines = False
        sheet.row_dimensions[1].height = 30
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.font = Font(color="FFFFFF", bold=True, size=11)
            cell.alignment = Alignment(vertical="center")
        for col in sheet.columns:
            cells = list(col)
            width = min(52, max(14, max(len(str(c.value or "")) for c in cells[:100]) + 2))
            sheet.column_dimensions[get_column_letter(cells[0].column)].width = width
        for row in sheet.iter_rows(min_row=2):
            if row[0].row % 2 == 0:
                for cell in row:
                    cell.fill = PatternFill("solid", fgColor=PALE)
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.print_options.horizontalCentered = True
    book.save(path)
    return path
