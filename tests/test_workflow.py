import asyncio
import csv
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from property_intel.analytics import comparables, market_breakdown, overview
from property_intel.app import PropertyIntel
from property_intel.db import Property, RawObservation, Transaction, open_db
from property_intel.demo import seed
from property_intel.pipeline import (
    best_match,
    import_rows,
    match_score,
    normalise_address,
    normalise_company,
    parse_area,
    parse_date,
    parse_money,
    read_rows,
    resolve,
    suggest_mapping,
    validate,
)
from property_intel.report import export_report


def test_parsing_validation_and_false_merges(tmp_path: Path) -> None:
    assert parse_money("$4.2m") == parse_money("4,200,000") == 4_200_000
    assert parse_money("850k") == 850_000
    assert parse_area("0.125 ha") == parse_area("1,250 m²") == 1250
    assert parse_date("14 Mar 2026") == date(2026, 3, 14)
    assert normalise_address("123 QUEEN STREET.") == "123 queen st"
    assert normalise_company("Example Holdings Limited") == normalise_company(
        "EXAMPLE HOLDINGS LTD."
    )
    for bad in ("03/04/2026", "not a date"):
        try:
            parse_date(bad)
        except ValueError:
            pass
        else:
            assert False, bad
    property = Property(canonical_address="123 Queen St", sector="QSR")
    assert match_score({"address": "123 Queen Street"}, property)[0] >= 0.93
    assert match_score({"address": "12 Queen Street"}, property)[0] == 0
    property.canonical_address = "Unit 2, 40 Example Rd"
    assert match_score({"address": "Unit 1, 40 Example Rd"}, property)[0] == 0
    issues = validate(
        {
            "sale_price": -1,
            "land_area_m2": 0,
            "transaction_date": date.today() + timedelta(days=1),
            "reported_yield": 0.25,
        }
    )
    assert {issue.field for issue in issues if issue.severity == "error"} >= {
        "sale_price",
        "land_area_m2",
        "transaction_date",
    }
    assert any(issue.field == "reported_yield" and issue.severity == "warning" for issue in issues)


def test_demo_analytics_comparables_excel(tmp_path: Path) -> None:
    session = open_db(tmp_path / "demo.db")
    seed(session)
    seed(session)
    stats = overview(session)
    assert (stats["properties"], stats["transactions"]) == (420, 610)
    assert sum(group["transactions"] for group in market_breakdown(session, "quarter")) == 610
    assert sum(group["transactions"] for group in market_breakdown(session, "region")) == 610
    assert len(market_breakdown(session, "sector")) == 3
    assert session.scalar(select(func.count(RawObservation.id))) == 760
    review_scores = session.scalars(
        select(RawObservation.match_score).where(RawObservation.status == "REVIEW")
    ).all()
    assert len(review_scores) == 12 and all(0.75 <= score < 0.93 for score in review_scores)
    property = session.scalars(select(Property)).first()
    assert property is not None
    ranked = comparables(session, property)
    assert len(ranked) == 5
    assert ranked == sorted(ranked, key=lambda item: item[1], reverse=True)
    path = export_report(session, tmp_path / "report.xlsx", property)
    workbook = load_workbook(path)
    assert workbook.sheetnames == [
        "01 Executive Summary",
        "02 Transactions",
        "03 Comparable Sales",
        "04 Sector Analysis",
        "05 Data Quality",
    ]
    assert workbook["02 Transactions"].max_row == 611
    assert workbook["02 Transactions"].freeze_panes == "C2"
    assert "TransactionSales" in workbook["02 Transactions"].tables
    assert workbook["03 Comparable Sales"].max_row == 11


def test_import_and_review(tmp_path: Path) -> None:
    session = open_db(tmp_path / "import.db")
    existing = Property(
        canonical_address="123 Queen St, Auckland",
        sector="SERVICE_STATION",
        city="Auckland",
        land_area_m2=1250,
        tenant_name="BP",
    )
    session.add(existing)
    session.commit()
    path = tmp_path / "sales.csv"
    with path.open("w", newline="") as file:
        writer = csv.DictWriter(
            file,
            fieldnames=[
                "Property Address",
                "Settlement Date",
                "Sale Amount",
                "Site Area",
                "Tenant",
            ],
        )
        writer.writeheader()
        writer.writerow(
            {
                "Property Address": "123 Queen Street, Auckland",
                "Settlement Date": "14 Mar 2026",
                "Sale Amount": "$4.2m",
                "Site Area": "1250 m2",
                "Tenant": "BP",
            }
        )
        writer.writerow(
            {
                "Property Address": "12 Queen St, Auckland",
                "Settlement Date": "2026-03-14",
                "Sale Amount": "$850k",
                "Site Area": "600 sqm",
                "Tenant": "BP",
            }
        )
        writer.writerow(
            {
                "Property Address": "123 Queen St, Auckland",
                "Settlement Date": "03/04/2026",
                "Sale Amount": "-2",
                "Site Area": "1250 sqm",
                "Tenant": "BP",
            }
        )
    rows = read_rows(path)
    mapping = suggest_mapping(list(rows[0]))
    job = import_rows(session, rows, mapping, path.name)
    assert (job.records_received, job.records_imported, job.records_rejected) == (3, 2, 1)
    assert session.scalar(select(func.count(Property.id))) == 2
    assert session.scalar(select(func.count(Transaction.id))) == 2
    assert best_match(session, {"address": "12 Queen St, Auckland"})[
        0
    ].canonical_address.startswith("12 ")

    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["Property Address", "Settlement Date", "Sale Amount", "Site Area", "Tenant"])
    sheet.append(["123 Queen Street, Auckland", "2026-03-14", "$4.2m", "1250 sqm", "Shell"])
    xlsx = tmp_path / "review.xlsx"
    workbook.save(xlsx)
    review_rows = read_rows(xlsx)
    review_job = import_rows(session, review_rows, suggest_mapping(list(review_rows[0])), xlsx.name)
    assert review_job.records_flagged == 1
    observation = session.scalars(
        select(RawObservation).where(RawObservation.status == "REVIEW")
    ).one()
    resolve(session, observation, merge=False)
    assert observation.status == "SEPARATE"
    assert session.scalar(select(func.count(Property.id))) == 3
    assert session.scalar(select(func.count(Transaction.id))) == 3


def test_tui_navigation(tmp_path: Path) -> None:
    session = open_db(tmp_path / "ui.db")
    seed(session)

    async def check() -> None:
        app = PropertyIntel(session, demo=True)
        async with app.run_test(size=(120, 35)) as pilot:
            assert app.focused and app.focused.id == "table"
            assert app.page == "overview" and len(app.rows) == 12
            await pilot.press("2")
            assert app.page == "transactions" and len(app.rows) == 610
            await pilot.press("3")
            assert app.page == "properties" and len(app.rows) == 420
            await pilot.press("4")
            assert app.page == "quality" and len(app.rows) == 37
            await pilot.press("5")
            assert app.page == "import"
            await pilot.press("escape", "/")
            assert app.page == "search"
            await pilot.press("escape", "3", "f")
            assert app.page == "properties"
            assert app.query_one("#advanced").display

    asyncio.run(check())
